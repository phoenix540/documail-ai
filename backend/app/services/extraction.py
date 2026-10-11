import io

from docx import Document as DocxDocument
from pypdf import PdfReader

SUPPORTED = (".pdf", ".docx", ".txt", ".md")


class UnsupportedFile(ValueError):
    pass


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
    elif name.endswith(".docx"):
        doc = DocxDocument(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    elif name.endswith((".txt", ".md")):
        text = data.decode("utf-8", errors="replace")
    else:
        raise UnsupportedFile(f"Unsupported file type. Use one of: {', '.join(SUPPORTED)}")
    return text.strip()
